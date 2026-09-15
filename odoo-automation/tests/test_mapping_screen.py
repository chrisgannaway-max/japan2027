"""Chart of accounts, the needs-mapping screen, and importing the filled worksheet."""
import importlib

import pytest

from conftest import FIXTURES
from pms_to_odoo.mapping import GLMapping
from pms_to_odoo.mapping_edit import escape_label, insert_rules, rule_for, suggest_account

CFG = FIXTURES.parent.parent / "config"
FULL_PEP = (CFG / "gl_mapping" / "hilton_pep.example.yaml").read_text()
# drop the rule that catches GUEST ROOM / GUEST ROOM HONORS, leaving those lines unmapped
PARTIAL_PEP = "\n".join(l for l in FULL_PEP.splitlines() if "^guest room( honors)?$" not in l) + "\n"


# ------------------------------------------------------------------ rule writing
def test_escape_label_is_readable_and_exact():
    import re
    for label in ["GUEST ROOM", "Sales Tax 4.125% Ex", "reservation_in_house(Offset)", "A/R Cash", "BEER + WINE"]:
        esc = escape_label(label)
        assert re.fullmatch(esc, label), f"{esc!r} should match {label!r}"
        assert "\\ " not in esc, "spaces should not be escaped"
    assert escape_label("Sales Tax 4.125% Ex") == r"Sales Tax 4\.125% Ex"


def test_rule_for_prefers_the_pms_code():
    assert rule_for("Room Charge", "RM", "revenue", "4000") == "  - {code: 'RM', account: '4000', section: revenue}"
    assert rule_for("GUEST ROOM", "", "revenue", "4000") == "  - {match: '^GUEST ROOM$', account: '4000', section: revenue}"
    assert "''" in rule_for("Joe's Fee", "", "revenue", "4020")          # quote doubled for YAML


def test_insert_rules_goes_to_the_top_and_keeps_comments():
    before = "# a comment\njournal: NA\nrules:\n  - {match: '.*', account: '9999'}   # catch-all\nignore:\n  - '^total'\n"
    after = insert_rules(before, [rule_for("X", "", "revenue", "4000")], "# note")
    assert "# a comment" in after and "# catch-all" in after and "ignore:" in after
    lines = after.splitlines()
    i_new, i_catch = lines.index("  - {match: '^X$', account: '4000', section: revenue}"), next(
        n for n, l in enumerate(lines) if "catch-all" in l)
    assert i_new < i_catch, "a new rule must win over an existing catch-all"
    import tempfile, os          # and the result is still valid YAML that GLMapping accepts
    with tempfile.NamedTemporaryFile("w", suffix=".yaml", delete=False) as fh:
        fh.write(after); tmp = fh.name
    try:
        assert len(GLMapping.load(tmp).rules) == 2
    finally:
        os.unlink(tmp)


def test_insert_rules_creates_the_block_when_missing():
    out = insert_rules("journal: NA\n", [rule_for("Y", "", "tax", "2200")])
    assert "rules:" in out and "'^Y$'" in out


def test_suggest_account_prefers_the_same_pms(tmp_path):
    props = {"A": {"pms": "PEP"}, "B": {"pms": "OPERA"}, "C": {"pms": "PEP"}}
    import tempfile, os
    def loader(prop):
        text = {"PEP": "journal: NA\nrules:\n  - {match: '^GUEST ROOM$', account: '4000'}\n",
                "OPERA": "journal: NA\nrules:\n  - {match: '^GUEST ROOM$', account: '9999'}\n"}[prop["pms"]]
        with tempfile.NamedTemporaryFile("w", suffix=".yaml", delete=False) as fh:
            fh.write(text); t = fh.name
        try:
            return GLMapping.load(t)
        finally:
            os.unlink(t)
    s = suggest_account({"label": "GUEST ROOM", "code": "", "section": "revenue"}, "PEP", props, loader, exclude="A")
    assert s.account == "4000" and "same PMS" in s.source
    s2 = suggest_account({"label": "NOTHING LIKE THIS", "code": "", "section": "revenue"}, "PEP", props, loader)
    assert not s2


# ------------------------------------------------------------------ portal
@pytest.fixture()
def client(tmp_path, monkeypatch):
    (tmp_path / "full.yaml").write_text(FULL_PEP)
    (tmp_path / "partial.yaml").write_text(PARTIAL_PEP)
    (tmp_path / "properties.yaml").write_text(
        "properties:\n"
        "  - {code: DONE1, name: Mapped sibling, pms: PEP, pms_property_id: XXXX, gl_mapping: full.yaml}\n"
        "  - {code: NEW1, name: New hotel, pms: PEP, pms_property_id: OKCON, gl_mapping: partial.yaml}\n")
    monkeypatch.setenv("PORTAL_DATA", str(tmp_path / "data"))
    monkeypatch.setenv("PORTAL_CONFIG", str(tmp_path / "properties.yaml"))
    monkeypatch.setenv("PORTAL_USERS", str(CFG / "users.example.yaml"))
    monkeypatch.setenv("PORTAL_SECRET", "test-secret")
    monkeypatch.setenv("PORTAL_STORE", "yaml")
    monkeypatch.delenv("ODOO_URL", raising=False)
    import portal.app as app_module
    importlib.reload(app_module)
    from fastapi.testclient import TestClient
    c = TestClient(app_module.app, follow_redirects=False)
    assert c.post("/login", data={"username": "admin", "password": "admin"}).status_code == 303
    return c, app_module, tmp_path


def load_accounts(client):
    csv = "code,name,type\n4000,Room revenue,income\n2200,Tax payable,liability\n1120,Card clearing,asset\n"
    r = client.post("/admin/accounts/upload", files={"file": ("coa.csv", csv.encode(), "text/csv")})
    assert r.status_code == 303
    return r


def test_chart_of_accounts_upload_and_search(client):
    c, app_module, _ = client
    assert "No accounts loaded yet" in c.get("/admin/accounts").text
    load_accounts(c)
    page = c.get("/admin/accounts").text
    assert "3 accounts loaded" in page and "Room revenue" in page
    assert "4000" in c.get("/admin/accounts?q=Room").text
    assert "1120" not in c.get("/admin/accounts?q=Room").text
    assert app_module.state.db.accounts_info()["count"] == 3


def test_map_missing_lines_then_rerun(client):
    c, app_module, tmp_path = client
    load_accounts(c)
    with open(FIXTURES / "pep_final_audit.txt", "rb") as fh:
        r = c.post("/upload", data={"property_code": "NEW1"}, files=[("files", ("audit.txt", fh, "text/plain"))])
    assert r.status_code == 200 and "Needs mapping" in r.text
    run_id = int(r.text.split("/runs/")[1].split('"')[0])
    # the run page offers the mapping screen
    assert "Map the missing lines" in c.get(f"/runs/{run_id}").text
    page = c.get(f"/admin/mapping/{run_id}").text
    assert "GUEST ROOM" in page and "GUEST ROOM HONORS" in page
    # the sibling property on the same PMS supplies the suggestion
    assert "DONE1 (same PMS)" in page and 'value="4000" selected' in page
    res = app_module.RunResult.from_json(app_module.state.db.get_run(run_id)["result_json"])
    form = {f"account_{i}": "4000" for i in range(len(res.unmapped))}
    r = c.post(f"/admin/mapping/{run_id}", data=form)
    assert r.status_code == 303 and "/runs/" in r.headers["location"]
    new_id = int(r.headers["location"].rsplit("/", 1)[1])
    detail = c.get(f"/runs/{new_id}").text
    assert "Balanced, ready to post" in detail and "92,570.78" in detail
    # the rule was written to the property's own mapping file, at the top, and is reusable
    text = (tmp_path / "partial.yaml").read_text()
    assert "added from the portal" in text and "'^GUEST ROOM$'" in text
    assert text.index("'^GUEST ROOM$'") < text.index("no show|late cancel")
    assert len(GLMapping.load(tmp_path / "partial.yaml").rules) == len(GLMapping.load(tmp_path / "full.yaml").rules) + 1


def test_ignore_writes_an_ignore_pattern(client):
    c, app_module, tmp_path = client
    load_accounts(c)
    with open(FIXTURES / "pep_final_audit.txt", "rb") as fh:
        r = c.post("/upload", data={"property_code": "NEW1"}, files=[("files", ("audit.txt", fh, "text/plain"))])
    run_id = int(r.text.split("/runs/")[1].split('"')[0])
    res = app_module.RunResult.from_json(app_module.state.db.get_run(run_id)["result_json"])
    form = {f"account_{i}": "__ignore__" for i in range(len(res.unmapped))}
    assert c.post(f"/admin/mapping/{run_id}", data=form).status_code == 303
    text = (tmp_path / "partial.yaml").read_text()
    assert "'^GUEST ROOM$'" in text and "ignore:" in text
    assert GLMapping.load(tmp_path / "partial.yaml").is_ignored(
        __import__("pms_to_odoo.models", fromlist=["ReportLine"]).ReportLine("GUEST ROOM", 1, "revenue"))


def test_worksheet_import(client):
    c, app_module, tmp_path = client
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Hilton PEP"
    for col, head in enumerate(["Section", "PMS code", "Line as printed on the report", "Amount on sample day",
                                "Activity on sample day", "Our suggestion (placeholder)", "What we think it is",
                                "YOUR ACCOUNT CODE", "ANSWER / notes"], 1):
        ws.cell(4, col, head)
    ws.cell(5, 1, "revenue"); ws.cell(5, 3, "GUEST ROOM"); ws.cell(5, 8, "4000")
    ws.cell(6, 1, "revenue"); ws.cell(6, 3, "GUEST ROOM HONORS"); ws.cell(6, 8, "4000")
    ws.cell(7, 1, "revenue"); ws.cell(7, 3, "NOT FILLED IN")           # no account -> skipped
    skip = wb.create_sheet("Start here"); skip["A1"] = "instructions"
    import io
    buf = io.BytesIO(); wb.save(buf)
    r = c.post("/admin/mapping/import/worksheet", files={"file": ("ws.xlsx", buf.getvalue(),
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert "DONE1: 2 rules" in r.text and "NEW1: 2 rules" in r.text      # both PEP properties
    assert "Start here" in r.text and "not a PMS tab" in r.text
    text = (tmp_path / "partial.yaml").read_text()
    assert "imported from the mapping worksheet" in text and "'^GUEST ROOM HONORS$'" in text
    # and the report now balances without touching anything else
    with open(FIXTURES / "pep_final_audit.txt", "rb") as fh:
        r2 = c.post("/upload", data={"property_code": "NEW1"}, files=[("files", ("audit.txt", fh, "text/plain"))])
    assert "Balanced, ready to post" in r2.text
