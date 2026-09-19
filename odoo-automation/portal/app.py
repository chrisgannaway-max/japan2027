"""FastAPI application.

    uvicorn portal.app:app --host 0.0.0.0 --port 8000        (or: python -m portal serve)

Environment:
    PORTAL_CONFIG   config/properties.yaml (default: config/properties.yaml or the example)
    PORTAL_USERS    config/users.yaml
    PORTAL_DATA     data/  (uploads + sqlite)
    PORTAL_SECRET   cookie signing secret (set one in production)
    ODOO_*          optional; enables "Post to Odoo" on the review screen
"""
from __future__ import annotations

import os
import re
import secrets
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Optional
from urllib.parse import quote_plus

from fastapi import BackgroundTasks, Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, PlainTextResponse, RedirectResponse, Response
from fastapi.templating import Jinja2Templates

from pms_to_odoo.export import bills_to_odoo_csv, entries_to_flat_csv, entries_to_odoo_csv
from pms_to_odoo.invoices import InvoiceData, InvoiceLine, create_vendor_bill, extract_invoice_auto, reader_in_use
from pms_to_odoo.invoices.to_odoo import load_expense_map
from pms_to_odoo.invoices.accounts import AccountAssigner
from pms_to_odoo.journal import format_entry, post_entry
from pms_to_odoo import odoo_check
from pms_to_odoo.odoo_client import OdooClient, OdooError, OdooSettings
from pms_to_odoo.mapping import GLMapping
from pms_to_odoo.mapping_edit import insert_rules, rule_for, suggest_account
from pms_to_odoo.invoices.sniff import looks_like_invoice
from pms_to_odoo.parsers import detect_pms as _detect_pms, get_parser, read_text as _read_text
from pms_to_odoo.pipeline import STATUS_LABELS, RunResult, load_properties, process_file, resolve

from . import mail
from .auth import (LoginThrottle, SessionSigner, User, UserStore, hash_password,
                    new_totp_secret, totp_qr_svg, totp_uri, unusable_password_hash, verify_totp)
from . import clock, daily, intake, poster, properties_import, scheduler, storage, users_import
from .db import Database
from .store import ConfigStore, PROPERTY_COLUMNS, read_mapping_text, store_mode, write_mapping_text

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
CONFIG = Path(os.environ.get("PORTAL_CONFIG", ROOT / "config" / "properties.yaml"))
USERS = Path(os.environ.get("PORTAL_USERS", ROOT / "config" / "users.yaml"))
DATA = Path(os.environ.get("PORTAL_DATA", ROOT / "data"))
ALLOWED_SUFFIXES = (".pdf", ".csv", ".xlsx", ".xlsm", ".eml", ".txt")
INVOICE_SUFFIXES = (".pdf", ".png", ".jpg", ".jpeg", ".txt")
VENDOR_TEMPLATES = ROOT / "config" / "vendor_templates.yaml"
EXPENSE_MAP = ROOT / "config" / "expense_categories.yaml"
VENDOR_ACCOUNTS = ROOT / "config" / "vendor_accounts.yaml"

app = FastAPI(title="Night Audit to Odoo", docs_url=None, redoc_url=None)
templates = Jinja2Templates(directory=str(HERE / "templates"))
templates.env.globals["STATUS_LABELS"] = STATUS_LABELS
# Stamps are stored in UTC and read by people in Oklahoma; the templates print them with
# `| when` (or `| when_short`) so no page has to remember that.
templates.env.filters["when"] = clock.show
templates.env.filters["when_short"] = clock.show_short
templates.env.filters["at"] = clock.show_local     # a time already on the hotels' clock
templates.env.globals["TZ_NAME"] = str(clock.TZ)


AUTOPOST_KEY = "ODOO_AUTOPOST"


def _on(name: str, default: str) -> bool:
    return os.environ.get(name, default).lower() in ("1", "on", "yes", "true")


#: Hotel logins and the upload page.  Off where packs arrive by e-mail and only the office signs
#: in.  Hidden rather than removed: a group that later wants its GMs to see their own nights
#: turns one variable back on, which is far cheaper than building it a second time.
HOTEL_UPLOADS = _on("HOTEL_UPLOADS", "on")

#: Vendor invoices.  A separate piece of work from the night audit, so a deployment doing only
#: night audits should not show a page nobody is meant to use.
INVOICES = _on("INVOICES", "off")


def _require_uploads() -> None:
    if not HOTEL_UPLOADS:
        raise HTTPException(404, "Uploads are off here: night-audit packs arrive by e-mail.")


def _require_invoices() -> None:
    if not INVOICES:
        raise HTTPException(404, "Invoice processing is not enabled here.")


def _after_intake() -> None:
    """Everything that should happen once a pack has been taken in, after the response has gone
    out: a hotel's pack arriving must not wait on Odoo, and Odoo being slow must not make a mail
    provider decide the delivery failed and send it again."""
    try:
        sent, why = daily.nudge_unmapped(state.db, state.props, base_url=mail.base_url())
        if sent:
            print("[nudge] told the office about nights needing account codes")
    except Exception as e:                       # noqa: BLE001
        print(f"[nudge] failed: {e}")
    if state.delivery != "odoo" or not state.odoo_enabled:
        return                                   # download mode: the queue simply waits
    try:
        print(f"[post] {poster.post_due(state.db, autopost=state.autopost)}")
    except Exception as e:                       # noqa: BLE001 - a background task must not die quietly
        print(f"[post] queue run failed: {e}")


def _users_file() -> Path:
    """config/users.yaml, or the example beside it on a fresh checkout."""
    return USERS if USERS.exists() else USERS.with_name("users.example.yaml")


class State:
    def __init__(self):
        self.mode = store_mode()                                   # "yaml" | "db"
        self.delivery = os.environ.get("DELIVERY_MODE", "download").lower()   # "download" | "odoo"
        self.db = Database(DATA / "portal.db")
        self.storage = storage.build(DATA)
        self.store = ConfigStore(DATA / "portal.db", DATA) if self.mode == "db" else None
        self.signer = SessionSigner()
        self.pending = SessionSigner(max_age=300)     # the 5 minutes between password and MFA code
        self.throttle = LoginThrottle()
        self.odoo_transport = os.environ.get("ODOO_TRANSPORT", "json2").lower()
        #: a stand-in Odoo, for rehearsing the whole path before the client has a server.
        #: Everything downstream treats it as configured, and every screen says DEMO.
        self.odoo_demo = self.odoo_transport == "demo"
        self.odoo_enabled = self.odoo_demo or bool(
            os.environ.get("ODOO_URL") and os.environ.get("ODOO_API_KEY"))
        self.reload_config()
        if self.autopost:
            print("[portal] entries will be POSTED in Odoo on arrival, not left as drafts")
        print(f"[portal] config from {self.mode}; database: {self.db.describe()}; "
              f"files: {self.storage.describe()}")
        print(f"[portal] cut-offs and times on screen are {clock.TZ}; it is "
              f"{clock.now():%H:%M} there now")
        if self.odoo_demo:
            print("[portal] ODOO_TRANSPORT=demo: entries go to a stand-in, not a real Odoo")

    @property
    def autopost_locked(self) -> bool:
        """Pinned on the host, so it cannot be flipped from the browser."""
        return bool(os.environ.get(AUTOPOST_KEY))

    def reload_config(self) -> None:
        stored = self.db.settings()
        mail.set_stored(stored)
        # Draft unless somebody has deliberately said otherwise.  A draft entry can be read and
        # deleted; a posted one needs a reversing entry, which is permanent in the ledger.
        self.autopost = (os.environ.get(AUTOPOST_KEY) or stored.get(AUTOPOST_KEY) or "no").lower() \
            in ("1", "yes", "true", "on")
        if self.store is not None:
            if not self.store.users():
                # First start against an empty database.  Without this nobody can sign in: the
                # logins live in the database, and the button that fills it is behind the login.
                counts = self.store.import_from_yaml(CONFIG, _users_file())
                print(f"[portal] empty database: seeded {counts['properties']} properties and "
                      f"{counts['users']} users from the config files -- change the passwords now")
            self.props = self.store.properties()
            self.users = UserStore(records=self.store.users())
        else:
            self.props = load_properties(CONFIG)
            self.users = UserStore(_users_file())


state = State()
scheduler.start(state)


# ------------------------------------------------------------------ auth helpers
def current_user(request: Request) -> Optional[User]:
    username = state.signer.verify(request.cookies.get("session"))
    return state.users.get(username) if username else None


def _https(request: Request) -> bool:
    """Whether to mark cookies Secure.  Honours X-Forwarded-Proto, since the site normally
    sits behind a proxy that terminates TLS."""
    mode = os.environ.get("PORTAL_SECURE_COOKIES", "auto").lower()
    if mode in ("always", "1", "true"):
        return True
    if mode in ("never", "0", "false"):
        return False
    proto = (request.headers.get("x-forwarded-proto") or "").split(",")[0].strip().lower()
    return (proto or request.url.scheme) == "https"


def _set_session(resp, request: Request, username: str) -> None:
    resp.set_cookie("session", state.signer.sign(username), httponly=True, samesite="lax",
                    secure=_https(request), max_age=state.signer.max_age)


def _client_ip(request: Request) -> str:
    fwd = (request.headers.get("x-forwarded-for") or "").split(",")[0].strip()
    return fwd or (request.client.host if request.client else "unknown")


def require_user(request: Request) -> User:
    u = current_user(request)
    if not u:
        raise HTTPException(status_code=303, headers={"Location": "/login"})
    # a role that must use MFA is sent to set it up before it can do anything else. Only
    # when it can actually be set up: enforcing it in file mode would lock the person out,
    # since the secret has nowhere to be stored.
    if (u.mfa_required() and state.store is not None and not u.needs_mfa()
            and not request.url.path.startswith(("/account", "/logout"))):
        raise HTTPException(status_code=303, headers={"Location": "/account?setup=1"})
    return u


def require_admin(request: Request) -> User:
    u = require_user(request)
    if not u.is_admin:
        raise HTTPException(status_code=403, detail="Admin only")
    return u


def visible_properties(u: User) -> dict[str, dict]:
    return state.props if u.is_admin else {c: p for c, p in state.props.items() if c in u.properties}


def render(request: Request, name: str, **ctx) -> HTMLResponse:
    ctx.setdefault("user", current_user(request))
    ctx.setdefault("odoo_enabled", state.odoo_enabled)
    ctx.setdefault("odoo_demo", state.odoo_demo)
    ctx.setdefault("invoice_reader", reader_in_use())
    ctx.setdefault("store_mode", state.mode)
    ctx.setdefault("delivery", state.delivery)
    ctx.setdefault("hotel_uploads", HOTEL_UPLOADS)
    ctx.setdefault("invoices_on", INVOICES)
    return templates.TemplateResponse(request, name, ctx)


@app.exception_handler(HTTPException)
async def redirect_on_303(request: Request, exc: HTTPException):
    if exc.status_code == 303:
        return RedirectResponse(exc.headers["Location"], status_code=303)
    return PlainTextResponse(str(exc.detail), status_code=exc.status_code)


# ------------------------------------------------------------------ login
@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request):
    return render(request, "login.html", error=None)


@app.post("/login")
def login(request: Request, username: str = Form(...), password: str = Form(...)):
    ip, username = _client_ip(request), username.strip()
    wait = state.throttle.locked_for(ip, username)
    if wait:
        return render(request, "login.html",
                      error=f"Too many attempts. Try again in {max(1, wait // 60)} minute(s).")
    u = state.users.authenticate(username, password)
    if not u:
        state.throttle.record_failure(ip, username)
        return render(request, "login.html", error="Wrong username or password.")
    state.throttle.record_success(ip, username)
    if u.needs_mfa():
        return render(request, "login_mfa.html", pending=state.pending.sign(u.username), error=None)
    resp = RedirectResponse("/" if u.is_admin else "/upload", status_code=303)
    _set_session(resp, request, u.username)
    return resp


@app.post("/login/mfa")
def login_mfa(request: Request, pending: str = Form(...), code: str = Form(...)):
    username = state.pending.verify(pending)
    if not username:
        return render(request, "login.html", error="That took too long. Please sign in again.")
    u = state.users.get(username)
    ip = _client_ip(request)
    if state.throttle.locked_for(ip, username):
        return render(request, "login.html", error="Too many attempts. Try again shortly.")
    if not u or not verify_totp(u.totp_secret, code):
        state.throttle.record_failure(ip, username)
        return render(request, "login_mfa.html", pending=pending, error="That code is not right.")
    state.throttle.record_success(ip, username)
    resp = RedirectResponse("/" if u.is_admin else "/upload", status_code=303)
    _set_session(resp, request, u.username)
    return resp


# ------------------------------------------------------------------ forgotten passwords
RESET_MINUTES = 60
#: Longer than a forgotten-password link: an imported account's first one has to survive
#: somebody getting to their e-mail the next morning.
SET_PASSWORD_MINUTES = 7 * 24 * 60
SAME_ANSWER = ("If that account exists and has an e-mail address on file, a reset link is on its way. "
               "The link lasts an hour.")


@app.get("/forgot", response_class=HTMLResponse)
def forgot_form(request: Request):
    return render(request, "forgot.html", message=None, error=None, can_email=mail.configured())


@app.post("/forgot", response_class=HTMLResponse)
def forgot(request: Request, who: str = Form(...)):
    u = state.users.find(who)
    if u and u.email and mail.configured():
        token = secrets.token_urlsafe(32)
        state.db.create_reset(u.username, token, RESET_MINUTES)
        link = f"{mail.base_url() or str(request.base_url).rstrip('/')}/reset?token={token}"
        subject, body = mail.reset_email(u.username, link, RESET_MINUTES)
        mail.send(u.email, subject, body)
    # the same answer either way, so the form cannot be used to discover who has an account
    return render(request, "forgot.html", message=SAME_ANSWER, error=None, can_email=mail.configured())


@app.get("/reset", response_class=HTMLResponse)
def reset_form(request: Request, token: str = ""):
    if not state.db.peek_reset(token):
        return render(request, "reset.html", token="", error="That link has expired or has already been used.")
    return render(request, "reset.html", token=token, error=None)


@app.post("/reset", response_class=HTMLResponse)
def reset(request: Request, token: str = Form(...), password: str = Form(...), confirm: str = Form(...)):
    if len(password) < 10:
        return render(request, "reset.html", token=token, error="Use at least 10 characters.")
    if password != confirm:
        return render(request, "reset.html", token=token, error="Those two do not match.")
    username = state.db.use_reset(token)
    if not username:
        return render(request, "reset.html", token="", error="That link has expired or has already been used.")
    if state.store is None:
        return render(request, "reset.html", token="",
                      error="Passwords are held in a file on this installation; ask an administrator to change it.")
    state.store.set_password(username, password)
    state.reload_config()
    return render(request, "reset.html", token="", done=True, error=None)


# ------------------------------------------------------------------ the user's own account
@app.get("/account", response_class=HTMLResponse)
def account(request: Request, setup: str = "", user: User = Depends(require_user)):
    secret = request.cookies.get("mfa_setup") or ""
    if not user.needs_mfa() and not secret:
        secret = new_totp_secret()
    ctx = dict(setup=bool(setup), qr=None, secret=secret, error=None)
    if secret and not user.needs_mfa():
        ctx["qr"] = totp_qr_svg(totp_uri(secret, user.username))
    resp = render(request, "account.html", **ctx)
    if secret and not user.needs_mfa():
        resp.set_cookie("mfa_setup", secret, httponly=True, samesite="lax", secure=_https(request), max_age=900)
    return resp


@app.post("/account/mfa/enable", response_class=HTMLResponse)
def mfa_enable(request: Request, code: str = Form(...), secret: str = Form(...), user: User = Depends(require_user)):
    if state.store is None:
        return render(request, "account.html", setup=False, qr=None, secret="",
                      error="MFA needs the database store (PORTAL_STORE=db).")
    if not verify_totp(secret, code):
        return render(request, "account.html", setup=True, secret=secret,
                      qr=totp_qr_svg(totp_uri(secret, user.username)), error="That code is not right. Try the next one.")
    state.store.set_mfa(user.username, secret, True)
    state.reload_config()
    resp = RedirectResponse("/account", status_code=303)
    resp.delete_cookie("mfa_setup")
    return resp


@app.post("/account/mfa/disable")
def mfa_disable(request: Request, user: User = Depends(require_user)):
    if user.mfa_required():
        raise HTTPException(400, "Your role must keep two-step sign-in switched on")
    if state.store is not None:
        state.store.set_mfa(user.username, "", False)
        state.reload_config()
    return RedirectResponse("/account", status_code=303)


@app.post("/admin/users/{username}/mfa-reset")
def admin_mfa_reset(username: str, user: User = Depends(require_admin)):
    """For the lost-phone case: an admin clears it and the person enrols again."""
    if state.store is None:
        raise HTTPException(400, "Needs the database store (PORTAL_STORE=db)")
    state.store.set_mfa(username, "", False)
    state.reload_config()
    return RedirectResponse(f"/admin?msg=Two-step+sign-in+cleared+for+{username}", status_code=303)


@app.get("/logout")
def logout():
    resp = RedirectResponse("/login", status_code=303)
    resp.delete_cookie("session")
    return resp


# ------------------------------------------------------------------ manager: upload
@app.get("/upload", response_class=HTMLResponse)
def upload_form(request: Request, user: User = Depends(require_user)):
    _require_uploads()
    props = visible_properties(user)
    runs = state.db.recent_runs(20, None if user.is_admin else list(props))
    return render(request, "upload.html", props=props, runs=runs, results=None)


@app.post("/upload", response_class=HTMLResponse)
async def upload(request: Request, background: BackgroundTasks, user: User = Depends(require_user),
                 property_code: str = Form(""), files: list[UploadFile] = File(...)):
    _require_uploads()
    props = visible_properties(user)
    allowed = None if user.is_admin else set(props)
    if property_code and property_code not in props:
        raise HTTPException(403, "Property not allowed")
    results: list[tuple[int, RunResult]] = []
    stamp = clock.now().strftime("%Y%m%d-%H%M%S")
    prefix = f"uploads/{property_code or 'unsorted'}/{stamp}"
    saved: list[tuple[Path, str]] = []
    for f in files:
        name = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(f.filename or "upload").name)
        if Path(name).suffix.lower() not in ALLOWED_SUFFIXES:
            results.append((0, RunResult(status="error", file_name=name, message="Only PDF, CSV, XLSX, EML or TXT files")))
            continue
        try:
            locator = state.storage.save(f"{prefix}/{name}", await f.read())
            saved.append((state.storage.local_path(locator), locator))
        except storage.StorageError as e:
            # A bucket that does not exist, or a key that is refused, is somebody's configuration
            # -- not a crash.  Say which file and why, on the page they are already looking at.
            results.append((0, RunResult(status="error", file_name=name, message=str(e))))
            print(f"[upload] storage refused {name}: {e}")
    consumed: set[Path] = set()
    for path, locator in saved:              # same folder, so SynXis pairs find each other
        if path in consumed:
            continue
        res = process_file(path, state.props, property_code or None, allowed)
        consumed.update(Path(c) for c in res.companions)
        run_id = state.db.add_run(uploaded_by=user.username, property_code=res.property_code,
                                  business_date=res.business_date.isoformat() if res.business_date else None,
                                  pms=res.pms, ref=res.ref, status=res.status, message=res.message,
                                  file_name=path.name, stored_path=locator, result_json=res.to_json())
        results.append((run_id, res))
    background.add_task(_after_intake)
    runs = state.db.recent_runs(20, None if user.is_admin else list(props))
    return render(request, "upload.html", props=props, runs=runs, results=results)


# ------------------------------------------------------------------ admin: dashboard
@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, day: Optional[str] = None, user: User = Depends(require_user)):
    if not user.is_admin:
        if not HOTEL_UPLOADS:
            return PlainTextResponse(
                "This site is for the accounting office. Night-audit packs are sent by e-mail, "
                "not uploaded here.", status_code=200)
        return RedirectResponse("/upload", status_code=303)
    d = date.fromisoformat(day) if day else clock.today() - timedelta(days=1)
    runs = state.db.runs_for_date(d.isoformat())
    latest: dict[str, object] = {}
    for r in runs:
        latest.setdefault(r["property_code"], r)
    rows = []
    for code, p in state.props.items():
        r = latest.get(code)
        rows.append({"code": code, "name": p.get("name", code), "pms": p.get("pms", ""), "run": r,
                     "result": RunResult.from_json(r["result_json"]) if r else None})
    ready = [r for r in runs if r["status"] == "ok"]
    waiting = [r for r in ready if not r["exported_at"] and not r["posted_at"]]
    already = [r for r in ready if r["exported_at"] or r["posted_at"]]
    return render(request, "dashboard.html", day=d, rows=rows, ready=ready, waiting=waiting, already=already,
                  prev_day=(d - timedelta(days=1)).isoformat(), next_day=(d + timedelta(days=1)).isoformat(),
                  dates=state.db.dates_with_runs())


@app.get("/runs/{run_id}", response_class=HTMLResponse)
def run_detail(request: Request, run_id: int, user: User = Depends(require_user)):
    r = state.db.get_run(run_id)
    if not r:
        raise HTTPException(404)
    if not user.is_admin and r["property_code"] not in user.properties:
        raise HTTPException(403)
    res = RunResult.from_json(r["result_json"])
    return render(request, "run.html", run=r, res=res,
                  entry_text=format_entry(res.entry) if res.entry else None)


@app.get("/runs/{run_id}/file")
def run_file(run_id: int, n: int = 0, user: User = Depends(require_user)):
    """The pack itself, back out of wherever it was filed.

    The portal keeps every uploaded file because it is the evidence behind an accounting
    entry, and until now there was no way to look at one again.  `n` picks the companion for
    a format that arrives in two halves; it indexes a list of locators the portal wrote
    itself, so nothing a browser sends ever reaches storage as a path.
    """
    import mimetypes

    r = state.db.get_run(run_id)
    if not r:
        raise HTTPException(404)
    if not user.is_admin and r["property_code"] not in user.properties:
        raise HTTPException(403)
    res = RunResult.from_json(r["result_json"])
    locators = [r["stored_path"]] + list(res.companions or [])
    if not 0 <= n < len(locators):
        raise HTTPException(404)
    name = r["file_name"] if n == 0 else Path(locators[n]).name
    try:
        body = state.storage.read(locators[n])
    except storage.StorageError as e:
        raise HTTPException(503, f"The file could not be fetched: {e}")
    except (FileNotFoundError, OSError):
        raise HTTPException(404, "That file is no longer in storage.")
    return Response(body, media_type=mimetypes.guess_type(name)[0] or "application/octet-stream",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.post("/runs/{run_id}/approve")
def approve(run_id: int, user: User = Depends(require_admin)):
    r = state.db.get_run(run_id)
    if not r or r["status"] != "ok":
        raise HTTPException(400, "Only balanced runs can be approved")
    state.db.mark(run_id, approved_at=clock.stamp(), approved_by=user.username)
    return RedirectResponse(f"/runs/{run_id}", status_code=303)


@app.post("/runs/{run_id}/post")
def post_to_odoo(run_id: int, user: User = Depends(require_admin), post_now: str = Form("no")):
    r = state.db.get_run(run_id)
    if not r or r["status"] != "ok":
        raise HTTPException(400, "Only balanced runs can be posted")
    if not state.odoo_enabled:
        raise HTTPException(400, "Odoo connection not configured (ODOO_URL / ODOO_API_KEY)")
    res = RunResult.from_json(r["result_json"])
    try:
        client = OdooClient.connect(OdooSettings.from_env())
        result = post_entry(res.entry, client, post=(post_now == "yes") or state.autopost)
    except OdooError as e:
        raise HTTPException(502, f"Odoo error: {e}") from e
    state.db.mark(run_id, posted_at=clock.stamp(), odoo_move_id=result.move_id)
    return RedirectResponse(f"/runs/{run_id}", status_code=303)


@app.post("/runs/{run_id}/reprocess")
def reprocess(run_id: int, user: User = Depends(require_admin)):
    r = state.db.get_run(run_id)
    if not r:
        raise HTTPException(404)
    state.reload_config()                            # pick up mapping edits
    res = process_file(state.storage.local_path(r["stored_path"]), state.props, r["property_code"] or None)
    new_id = state.db.add_run(uploaded_by=user.username, property_code=res.property_code,
                              business_date=res.business_date.isoformat() if res.business_date else None,
                              pms=res.pms, ref=res.ref, status=res.status, message=res.message,
                              file_name=r["file_name"], stored_path=r["stored_path"], result_json=res.to_json())
    return RedirectResponse(f"/runs/{new_id}", status_code=303)


# ------------------------------------------------------------------ admin: export
def _entries_for(day: str, only_approved: bool, include_exported: bool = False) -> list:
    """Balanced entries for a day.  Already-downloaded ones are left out unless asked for:
    handing the same entry out twice is how a day gets booked twice."""
    out = []
    for r in state.db.runs_for_date(day):
        if r["status"] != "ok" or (only_approved and not r["approved_at"]):
            continue
        if r["exported_at"] and not include_exported:
            continue
        res = RunResult.from_json(r["result_json"])
        if res.entry:
            out.append((r["id"], res.entry))
    return out


@app.get("/export/{day}.csv")
def export_day(day: str, fmt: str = "odoo", approved: str = "all", again: str = "no",
               user: User = Depends(require_admin)):
    date.fromisoformat(day)
    pairs = _entries_for(day, approved == "yes", include_exported=(again == "yes"))
    entries = [e for _, e in pairs]
    body = entries_to_odoo_csv(entries) if fmt == "odoo" else entries_to_flat_csv(entries)
    for run_id, _ in pairs:
        state.db.mark(run_id, exported_at=clock.stamp())
    return Response(body, media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="night-audit-{day}-{fmt}.csv"'})


# Not "/runs/{run_id}.csv": that is the same single path segment as "/runs/{run_id}", so whichever
# route is declared first wins and "2.csv" arrives at the HTML page as a run id that will not
# parse. A separate segment cannot collide, whatever order these end up in.
@app.get("/runs/{run_id}/csv")
def export_run(run_id: int, fmt: str = "odoo", user: User = Depends(require_admin)):
    r = state.db.get_run(run_id)
    if not r or r["status"] != "ok":
        raise HTTPException(404)
    res = RunResult.from_json(r["result_json"])
    body = entries_to_odoo_csv([res.entry]) if fmt == "odoo" else entries_to_flat_csv([res.entry])
    state.db.mark(run_id, exported_at=clock.stamp())
    return Response(body, media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="{res.ref}-{fmt}.csv"'})


# ------------------------------------------------------------------ mail arriving
INTAKE_TOKEN = os.environ.get("INTAKE_TOKEN", "")


@app.post("/intake/mail")
async def intake_mail(request: Request, background: BackgroundTasks):
    """An inbound mail provider POSTs one message here.

    The provider does not sign its requests, so the shared secret in INTAKE_TOKEN is what
    stands between this and anyone who guesses the URL.  Without it set, the route is off --
    an open endpoint that files attachments is not something to leave running by accident.
    """
    if not INTAKE_TOKEN:
        raise HTTPException(404)
    given = request.headers.get("X-Intake-Token") or request.query_params.get("token") or ""
    if not secrets.compare_digest(given, INTAKE_TOKEN):
        raise HTTPException(403, "bad or missing intake token")
    try:
        payload = await request.json()
    except Exception:  # noqa: BLE001
        raise HTTPException(400, "expected JSON") from None
    msg = intake.message_from_postmark(payload)
    try:
        res = intake.ingest(msg, db=state.db, storage=state.storage, props=state.props)
    except storage.StorageError as e:
        # 503 rather than 500: the provider should try this message again once somebody has
        # fixed the bucket, instead of giving up on a night that never reached us.
        print(f"[intake] storage refused {msg.message_id}: {e}")
        raise HTTPException(503, f"Could not store the attachment: {e}") from None
    print(f"[intake] {msg.sender or '?'} -> {len(res.created)} run(s), "
          f"{len(res.duplicates)} duplicate(s), {len(res.ignored)} ignored")
    # Always 200 once the message is ours: a report we could not parse is recorded as a run to
    # look at, not an error for the provider to retry until it gives up.
    background.add_task(_after_intake)
    return {"accepted": True,
            "runs": [{"id": rid, "property": r.property_code, "status": r.status}
                     for rid, r in res.created],
            "duplicates": res.duplicates, "ignored": res.ignored}


@app.get("/daily", response_class=PlainTextResponse)
def daily_report(day: Optional[str] = None, user: User = Depends(require_admin)):
    """The morning list, exactly as the e-mail sends it.  Plain text on purpose: it is the same
    thing read two ways, so there is no second version to drift."""
    d = date.fromisoformat(day) if day else scheduler.business_date_for(clock.now())
    return daily.as_text(daily.build(state.db, state.props, d), mail.base_url())


@app.post("/daily/send")
def daily_send(user: User = Depends(require_admin)):
    sent, why = scheduler.send_report(state, force=True)
    return RedirectResponse("/admin?" + ("msg=Report+sent" if sent else f"err={why}"), status_code=303)


@app.post("/admin/odoo/check", response_class=HTMLResponse)
def odoo_check_run(request: Request, user: User = Depends(require_admin)):
    """Ask Odoo everything a night will need, before a night depends on the answer.

    Creates nothing, so it is safe against a client's live server.
    """
    if not state.odoo_enabled:
        return RedirectResponse("/admin?err=No+Odoo+connection+configured", status_code=303)
    try:
        client = OdooClient.connect(OdooSettings.from_env())
    except OdooError as e:
        return render(request, "admin_odoo_check.html",
                      check=odoo_check.Check(reachable=False, why=str(e)), demo=state.odoo_demo)
    return render(request, "admin_odoo_check.html", check=odoo_check.run(client, state.props),
                  demo=state.odoo_demo)


@app.post("/admin/queue/run")
def queue_run_now(user: User = Depends(require_admin)):
    """Try the posting queue this minute rather than at the next pass.

    The same call the loop makes, so what happens here is what would have happened on its own.
    """
    if state.delivery != "odoo" or not state.odoo_enabled:
        return RedirectResponse("/admin?err=Set+DELIVERY_MODE%3Dodoo+with+an+Odoo+connection+first",
                                status_code=303)
    try:
        summary = poster.post_due(state.db, autopost=state.autopost)
    except Exception as e:                       # noqa: BLE001
        return RedirectResponse(f"/admin?err={quote_plus(f'Queue run failed: {e}')}", status_code=303)
    done, already, failed = len(summary.posted), len(summary.existing), len(summary.failed)
    if not (done or already or failed):
        return RedirectResponse("/admin?msg=Nothing+was+waiting+to+go", status_code=303)
    bits = ([f"{done} sent to Odoo"] if done else []) \
        + ([f"{already} already there"] if already else []) \
        + ([f"{failed} failed"] if failed else [])
    return RedirectResponse(f"/admin?msg={quote_plus(', '.join(bits))}", status_code=303)


@app.post("/admin/report-at")
def set_report_at(user: User = Depends(require_admin), report_at: str = Form("")):
    """The time the morning list goes out.  Blank goes back to the latest property cut-off."""
    raw = report_at.strip()
    if raw:
        try:
            hh, _, mm = raw.partition(":")
            if not (0 <= int(hh) <= 23 and 0 <= int(mm or 0) <= 59):
                raise ValueError
        except ValueError:
            return RedirectResponse("/admin?err=Give+a+time+as+HH%3AMM%2C+e.g.+06%3A30",
                                    status_code=303)
    if os.environ.get("REPORT_AT"):
        return RedirectResponse("/admin?err=REPORT_AT+is+pinned+on+the+host", status_code=303)
    state.db.save_settings({"REPORT_AT": raw}, user.username)
    mail.set_stored(state.db.settings())
    where = f"at {raw}" if raw else "back to the latest property cut-off"
    return RedirectResponse(f"/admin?msg={quote_plus('Morning list ' + where)}", status_code=303)


@app.get("/health", response_class=PlainTextResponse)
def health():
    return "ok"


# ------------------------------------------------------------------ invoices
def _expense_categories() -> dict[str, str]:
    path = EXPENSE_MAP if EXPENSE_MAP.exists() else EXPENSE_MAP.with_name("expense_categories.example.yaml")
    return load_expense_map(path)


def _assigner() -> AccountAssigner:
    path = VENDOR_ACCOUNTS if VENDOR_ACCOUNTS.exists() else VENDOR_ACCOUNTS.with_name("vendor_accounts.example.yaml")
    return AccountAssigner(path, state.db.remembered_account)


def _invoice_status(fields: dict, inv_id: int = 0) -> tuple[str, str]:
    """ready when the bill can go straight into the export; otherwise needs_review + why."""
    problems = []
    if not fields.get("vendor_name"):
        problems.append("vendor missing")
    if not fields.get("invoice_number"):
        problems.append("invoice number missing")
    if not fields.get("invoice_date"):
        problems.append("invoice date missing")
    try:
        if Decimal(str(fields.get("total") or 0)) <= 0:
            problems.append("total missing")
    except Exception:  # noqa: BLE001
        problems.append("total not a number")
    if not fields.get("account_code"):
        problems.append("no expense account could be assigned")
    if state.db.find_duplicate_invoice(fields.get("vendor_name", ""), fields.get("invoice_number", ""), inv_id):
        problems.append("possible duplicate")
    return ("ready", "") if not problems else ("needs_review", "; ".join(problems))


def _invoice_form_ctx(request: Request, user: User, **extra):
    props = visible_properties(user)
    invoices = state.db.list_invoices(None, None if user.is_admin else list(props), 30)
    return dict(props=props, invoices=invoices, categories=_expense_categories(), **extra)


@app.get("/invoices", response_class=HTMLResponse)
def invoices_page(request: Request, user: User = Depends(require_user)):
    _require_invoices()
    return render(request, "invoices.html", **_invoice_form_ctx(request, user, draft=None, inv_id=None))


def _create_invoice(path: Path, property_code: str, username: str, file_name: str,
                    locator: Optional[str] = None) -> int:
    """Read one invoice file and store it as a draft. Shared by the invoice page and the
    'this is actually an invoice' button on the night-audit page."""
    data, reader = extract_invoice_auto(path, VENDOR_TEMPLATES)
    description = data.lines[0].description if data.lines and data.lines[0].description != "Invoice total" else ""
    account, how = _assigner().assign(data.vendor_name, description or data.review_notes)
    fields = dict(property_code=property_code, vendor_name=data.vendor_name, vendor_tax_id=data.vendor_tax_id or "",
                  invoice_number=data.invoice_number, invoice_date=data.invoice_date, due_date=data.due_date or "",
                  subtotal=str(data.dec("subtotal")), tax_amount=str(data.dec("tax_amount")), total=str(data.dec("total")),
                  account_code=account or "", description=description or f"Invoice {data.invoice_number}".strip(),
                  notes=f"account by {how}; {data.review_notes}".strip("; "))
    status, why = _invoice_status(fields)
    if why:
        fields["notes"] = f"{why}; {fields['notes']}"
    inv_id = state.db.add_invoice(uploaded_by=username, file_name=file_name,
                                  stored_path=locator or str(path), reader=reader,
                                  confidence=data.confidence, **fields)
    state.db.update_invoice(inv_id, status=status)
    return inv_id


def _night_audit_report_in(path: Path) -> Optional[str]:
    """The PMS name if this file is a night-audit report, else None. Never raises."""
    if path.suffix.lower() not in (".pdf", ".txt", ".csv", ".eml", ".xlsx", ".xlsm"):
        return None
    try:
        return _detect_pms(_read_text(path))
    except Exception:  # noqa: BLE001 - an unreadable file is simply not a report
        return None


@app.post("/invoices/upload", response_class=HTMLResponse)
async def invoice_upload(request: Request, user: User = Depends(require_user),
                         property_code: str = Form(""), file: UploadFile = File(...)):
    _require_invoices()
    props = visible_properties(user)
    if property_code and property_code not in props:
        raise HTTPException(403, "Property not allowed")
    if not property_code and len(props) == 1:
        property_code = next(iter(props))
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(file.filename or "invoice").name)
    if Path(name).suffix.lower() not in INVOICE_SUFFIXES:
        return render(request, "invoices.html", **_invoice_form_ctx(request, user, draft=None, inv_id=None,
                      error="Only PDF, PNG, JPG or TXT invoices"))
    when = clock.now()
    key = (f"invoices/{property_code or 'unsorted'}/{when:%Y%m}/{when:%Y%m%d-%H%M%S}-{name}")
    try:
        locator = state.storage.save(key, await file.read())
        target = state.storage.local_path(locator)
    except storage.StorageError as e:
        return render(request, "invoices.html", **_invoice_form_ctx(request, user, draft=None,
                      inv_id=None, error=str(e)))
    pms = _night_audit_report_in(target)
    if pms:
        state.storage.delete(locator)
        return render(request, "invoices.html", **_invoice_form_ctx(request, user, draft=None, inv_id=None,
                      error=f"That is a {pms} night-audit report, not an invoice. Upload it on the Night audit page."))
    inv_id = _create_invoice(target, property_code, user.username, name, locator)
    return RedirectResponse(f"/invoices/{inv_id}", status_code=303)


@app.get("/invoices/{inv_id}", response_class=HTMLResponse)
def invoice_detail(request: Request, inv_id: int, user: User = Depends(require_user)):
    _require_invoices()
    inv = state.db.get_invoice(inv_id)
    if not inv:
        raise HTTPException(404)
    if not user.is_admin and inv["property_code"] not in user.properties:
        raise HTTPException(403)
    dup = state.db.find_duplicate_invoice(inv["vendor_name"], inv["invoice_number"], inv_id)
    return render(request, "invoices.html", **_invoice_form_ctx(request, user, draft=inv, inv_id=inv_id, duplicate=dup))


@app.post("/invoices/{inv_id}", response_class=HTMLResponse)
async def invoice_save(request: Request, inv_id: int, user: User = Depends(require_user)):
    _require_invoices()
    inv = state.db.get_invoice(inv_id)
    if not inv:
        raise HTTPException(404)
    if not user.is_admin and inv["property_code"] not in user.properties:
        raise HTTPException(403)
    form = await request.form()
    fields = {k: (form.get(k) or "").strip() for k in Database.INVOICE_FIELDS}
    if not user.is_admin and fields["property_code"] not in user.properties:
        raise HTTPException(403)
    for k in ("subtotal", "tax_amount", "total"):
        v = fields[k].replace("$", "").replace(",", "")
        fields[k] = str(Decimal(v)) if v else "0"
    if not fields["subtotal"] or fields["subtotal"] == "0":
        fields["subtotal"] = str(Decimal(fields["total"]) - Decimal(fields["tax_amount"]))
    action = form.get("action", "save")
    if not fields["account_code"]:
        fields["account_code"] = _assigner().assign(fields["vendor_name"], fields["description"])[0] or ""
    status, why = _invoice_status(fields, inv_id)
    if action == "hold" and user.is_admin:
        status = "held"
    elif action == "reject" and user.is_admin:
        status = "rejected"
    elif inv["status"] in ("exported", "posted"):
        status = inv["status"]                      # already out the door: edits do not un-export
    elif action == "release" and user.is_admin:
        status = "ready" if not why or why == "possible duplicate" else "needs_review"
    if why and status == "needs_review":
        fields["notes"] = why
    state.db.update_invoice(inv_id, status=status, **fields)
    return RedirectResponse(f"/invoices/{inv_id}", status_code=303)


@app.get("/invoices/export/bills.csv")
def invoices_export(status: str = "ready", user: User = Depends(require_admin)):
    _require_invoices()
    """Default: every ready bill not yet exported. Exporting remembers each vendor's account."""
    if status == "ready":
        rows = [r for r in state.db.list_invoices("ready", None, 1000)]
    elif status == "all":
        rows = state.db.list_invoices(None, None, 1000)
    else:
        rows = state.db.list_invoices(status, None, 1000)
    body = bills_to_odoo_csv([dict(r) for r in rows])
    now = clock.stamp()
    for r in rows:
        if r["status"] == "ready":
            state.db.update_invoice(r["id"], exported_at=now, status="exported")
            state.db.remember_account(r["vendor_name"], r["account_code"])
    return Response(body, media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="vendor-bills-{clock.now():%Y%m%d-%H%M}.csv"'})


@app.post("/invoices/{inv_id}/post")
def invoice_post(inv_id: int, user: User = Depends(require_admin)):
    _require_invoices()
    inv = state.db.get_invoice(inv_id)
    if not inv or inv["status"] not in ("ready", "exported"):
        raise HTTPException(400, "Bill is not ready (needs review, held or rejected)")
    if not state.odoo_enabled:
        raise HTTPException(400, "Odoo connection not configured (ODOO_URL / ODOO_API_KEY)")
    data = InvoiceData(vendor_name=inv["vendor_name"], vendor_tax_id=inv["vendor_tax_id"] or None,
                       invoice_number=inv["invoice_number"], invoice_date=inv["invoice_date"], due_date=inv["due_date"] or None,
                       subtotal=inv["subtotal"], tax_amount=inv["tax_amount"], total=inv["total"],
                       lines=[InvoiceLine(description=inv["description"] or f"Invoice {inv['invoice_number']}", quantity="1",
                                          unit_price=inv["subtotal"], amount=inv["subtotal"], category_hint="")],
                       confidence=inv["confidence"] or "medium", review_notes=inv["notes"] or "")
    prop = state.props.get(inv["property_code"], {})
    try:
        client = OdooClient.connect(OdooSettings.from_env())
        res = create_vendor_bill(data, client, str(state.storage.local_path(inv["stored_path"])),
                                 company=prop.get("company"),
                                 default_account=inv["account_code"] or None, create_missing_vendor=True)
    except OdooError as e:
        raise HTTPException(502, f"Odoo error: {e}") from e
    if res.move_id:
        state.db.update_invoice(inv_id, status="posted", posted_at=clock.stamp(),
                                odoo_move_id=res.move_id)
        state.db.remember_account(inv["vendor_name"], inv["account_code"])
    return RedirectResponse(f"/invoices/{inv_id}", status_code=303)


# ------------------------------------------------------------------ dashboard: send a whole day to Odoo
@app.post("/export/{day}/post")
def post_day(day: str, post_now: str = Form("no"), user: User = Depends(require_admin)):
    if not state.odoo_enabled:
        raise HTTPException(400, "Odoo connection not configured (ODOO_URL / ODOO_API_KEY)")
    client = OdooClient.connect(OdooSettings.from_env())
    for r in state.db.runs_for_date(day):
        if r["status"] != "ok" or r["posted_at"]:
            continue
        res = RunResult.from_json(r["result_json"])
        try:
            result = post_entry(res.entry, client, post=(post_now == "yes") or state.autopost)
        except OdooError as e:
            raise HTTPException(502, f"Odoo error on {res.ref}: {e}") from e
        state.db.mark(r["id"], posted_at=clock.stamp(), odoo_move_id=result.move_id)
    return RedirectResponse(f"/?day={day}", status_code=303)


# ------------------------------------------------------------------ admin: configuration (db store)
def require_db_store():
    if state.store is None:
        raise HTTPException(400, "Configuration is read from files (PORTAL_STORE=yaml). Set PORTAL_STORE=db to edit it here.")
    return state.store


@app.get("/admin", response_class=HTMLResponse)
def admin_home(request: Request, user: User = Depends(require_admin), msg: str = "", err: str = ""):
    ctx = dict(msg=msg, err=err, config_path=str(CONFIG), users_path=str(USERS),
               accounts_count=state.db.accounts_info()["count"],
               autopost=state.autopost, autopost_locked=state.autopost_locked,
               sched=scheduler.status(state), report_at=mail.setting("REPORT_AT"),
               report_at_locked=bool(os.environ.get("REPORT_AT")),
               queued=len(state.db.runs_awaiting_post()) if state.delivery == "odoo" else 0)
    if state.store is not None:
        ctx["properties"] = [state.store.get_property(c) | {"enabled": state.store.get_property(c)["enabled"]} for c in
                             state.store.properties(include_disabled=True)]
        ctx["users_list"] = state.store.all_users()
    else:
        ctx["properties"] = [{k: p.get(k, "") for k in PROPERTY_COLUMNS if k != "mapping_yaml"} | {"enabled": 1} for p in state.props.values()]
        ctx["users_list"] = [{"username": u.username, "role": u.role, "properties": ",".join(u.properties), "enabled": 1}
                             for u in state.users.users.values()]
    return render(request, "admin.html", **ctx)


@app.post("/admin/import-yaml")
def admin_import(user: User = Depends(require_admin), overwrite: str = Form("no")):
    store = require_db_store()
    counts = store.import_from_yaml(CONFIG, USERS, overwrite=(overwrite == "yes"))
    state.reload_config()
    return RedirectResponse(f"/admin?msg=Imported+{counts['properties']}+properties+and+{counts['users']}+users", status_code=303)


@app.get("/admin/properties/{code}", response_class=HTMLResponse)
def admin_property_form(request: Request, code: str, user: User = Depends(require_admin)):
    store = require_db_store()
    prop = store.get_property(code) if code != "new" else {k: "" for k in PROPERTY_COLUMNS} | {"enabled": 1}
    if prop is None:
        raise HTTPException(404)
    return render(request, "admin_property.html", prop=prop, is_new=(code == "new"), err="")


@app.post("/admin/properties/{code}", response_class=HTMLResponse)
async def admin_property_save(request: Request, code: str, user: User = Depends(require_admin)):
    store = require_db_store()
    form = await request.form()
    if form.get("action") == "delete" and code != "new":
        store.delete_property(code)
        state.reload_config()
        return RedirectResponse("/admin?msg=Property+deleted", status_code=303)
    fields = {k: (form.get(k) or "") for k in PROPERTY_COLUMNS}
    fields["enabled"] = "1" if form.get("enabled") else "0"
    try:
        store.save_property(**fields)
    except ValueError as e:
        return render(request, "admin_property.html", prop=fields, is_new=(code == "new"), err=str(e))
    state.reload_config()
    return RedirectResponse(f"/admin?msg=Saved+{fields['code']}", status_code=303)


@app.post("/admin/users")
async def admin_user_save(request: Request, user: User = Depends(require_admin)):
    store = require_db_store()
    form = await request.form()
    try:
        if form.get("action") == "delete":
            store.delete_user(form.get("username", ""))
        else:
            props = [p.strip() for p in (form.get("properties") or "").replace(";", ",").split(",")]
            store.save_user(form.get("username", ""), form.get("role", "manager"), props,
                            password=(form.get("password") or None),
                            # bool("0") is True, so the old reading of this could never switch a
                            # login off -- and users() only returns enabled=1, so "no" here
                            # is what stops somebody signing in.  Absent still means yes, because
                            # the "add a login" form below has no such field.
                            enabled=str(form.get("enabled", "1")).strip().lower()
                                    not in ("0", "no", "false", "off", ""),
                            email=(form.get("email") or ""))
    except ValueError as e:
        return RedirectResponse(f"/admin?err={e}", status_code=303)
    state.reload_config()
    return RedirectResponse("/admin?msg=User+saved", status_code=303)


# ------------------------------------------------------------------ chart of accounts
def _mapping_loader(prop: dict):
    """GLMapping for a property, however its mapping is stored (used for suggestions)."""
    import tempfile
    text = read_mapping_text(prop, state.store)
    if not text.strip():
        return None
    with tempfile.NamedTemporaryFile("w", suffix=".yaml", delete=False) as fh:
        fh.write(text); tmp = fh.name
    try:
        return GLMapping.load(tmp)
    finally:
        Path(tmp).unlink(missing_ok=True)


def _accounts() -> list:
    return state.db.list_accounts()


@app.get("/admin/accounts", response_class=HTMLResponse)
def accounts_page(request: Request, user: User = Depends(require_admin), q: str = "", msg: str = "", err: str = ""):
    return render(request, "admin_accounts.html", accounts=state.db.list_accounts(q, 500),
                  info=state.db.accounts_info(), q=q, msg=msg, err=err)


@app.post("/admin/accounts/upload")
async def accounts_upload(request: Request, user: User = Depends(require_admin), file: UploadFile = File(...)):
    import csv as _csv
    import io as _io
    raw = (await file.read()).decode("utf-8-sig", "replace")
    rows = list(_csv.DictReader(_io.StringIO(raw)))
    if not rows:
        return RedirectResponse("/admin/accounts?err=That+file+has+no+rows", status_code=303)
    keys = {k.strip().lower(): k for k in rows[0]}
    def pick(*names):
        for n in names:
            if n in keys: return keys[n]
        return None
    kc, kn, kt = pick("code", "account code", "id"), pick("name", "account name", "label"), pick("type", "account_type", "account type")
    if not kc:
        return RedirectResponse("/admin/accounts?err=No+'code'+column+found", status_code=303)
    n = state.db.replace_accounts([{"code": r.get(kc), "name": r.get(kn) if kn else "",
                                    "account_type": r.get(kt) if kt else ""} for r in rows], f"upload: {file.filename}")
    return RedirectResponse(f"/admin/accounts?msg=Loaded+{n}+accounts", status_code=303)


@app.post("/admin/accounts/pull")
def accounts_pull(user: User = Depends(require_admin), company: str = Form("")):
    if not state.odoo_enabled:
        return RedirectResponse("/admin/accounts?err=Odoo+connection+not+configured", status_code=303)
    try:
        client = OdooClient.connect(OdooSettings.from_env())
        rows = client.chart_of_accounts(client.company_id(company or None))
    except OdooError as e:
        return RedirectResponse(f"/admin/accounts?err=Odoo+error:+{e}", status_code=303)
    n = state.db.replace_accounts(rows, "pulled from Odoo")
    return RedirectResponse(f"/admin/accounts?msg=Pulled+{n}+accounts+from+Odoo", status_code=303)


# ------------------------------------------------------------------ needs mapping
@app.get("/admin/mapping/{run_id}", response_class=HTMLResponse)
def mapping_page(request: Request, run_id: int, user: User = Depends(require_admin), err: str = ""):
    r = state.db.get_run(run_id)
    if not r:
        raise HTTPException(404)
    res = RunResult.from_json(r["result_json"])
    prop = state.props.get(res.property_code)
    if prop is None:
        raise HTTPException(400, f"Property {res.property_code} is not configured")
    rows = []
    for line in res.unmapped:
        sug = suggest_account(line, res.pms, state.props, _mapping_loader, exclude=res.property_code)
        rows.append({**line, "suggested": sug.account, "suggested_from": sug.source})
    return render(request, "admin_mapping.html", run=r, res=res, rows=rows, accounts=_accounts(), err=err)


@app.post("/admin/mapping/{run_id}")
async def mapping_save(request: Request, run_id: int, user: User = Depends(require_admin)):
    r = state.db.get_run(run_id)
    if not r:
        raise HTTPException(404)
    res = RunResult.from_json(r["result_json"])
    prop = state.props.get(res.property_code)
    if prop is None:
        raise HTTPException(400, f"Property {res.property_code} is not configured")
    form = await request.form()
    new_rules, ignored = [], 0
    for i, line in enumerate(res.unmapped):
        choice = (form.get(f"account_{i}") or "").strip()
        if not choice:
            continue
        if choice == "__ignore__":
            ignored += 1
            continue
        rule = rule_for(line["label"], line.get("code", ""), line.get("section", ""), choice)
        if rule not in new_rules:          # the same label can appear on several report lines
            new_rules.append(rule)
    if not new_rules and not ignored:
        return RedirectResponse(f"/admin/mapping/{run_id}?err=Nothing+chosen", status_code=303)
    text = read_mapping_text(prop, state.store)
    if ignored:                       # "not an accounting line" -> an ignore pattern, not a rule
        from pms_to_odoo.mapping_edit import escape_label, yq
        pats = [f"  - {yq('^' + escape_label(l['label']) + '$')}"
                for i, l in enumerate(res.unmapped) if (form.get(f"account_{i}") or "") == "__ignore__"]
        if re.search(r"^ignore:", text, re.M):
            text = re.sub(r"^ignore:.*$", "ignore:\n" + "\n".join(pats), text, count=1, flags=re.M)
        else:
            text = text.rstrip("\n") + "\nignore:\n" + "\n".join(pats) + "\n"
    text = insert_rules(text, new_rules, f"# added from the portal {clock.today().isoformat()} by {user.username}")
    try:
        where = write_mapping_text(prop, state.store, text)
    except ValueError as e:
        return RedirectResponse(f"/admin/mapping/{run_id}?err={e}", status_code=303)
    state.reload_config()
    res2 = process_file(state.storage.local_path(r["stored_path"]), state.props, res.property_code)
    new_id = state.db.add_run(uploaded_by=f"{user.username} (remapped)", property_code=res2.property_code,
                              business_date=res2.business_date.isoformat() if res2.business_date else None,
                              pms=res2.pms, ref=res2.ref, status=res2.status, message=res2.message,
                              file_name=r["file_name"], stored_path=r["stored_path"], result_json=res2.to_json())
    print(f"[mapping] {len(new_rules)} rule(s), {ignored} ignore(s) for {res.property_code} -> {where}")
    return RedirectResponse(f"/runs/{new_id}" if res2.status == "ok" else f"/admin/mapping/{new_id}", status_code=303)


# ------------------------------------------------------------------ hotels and logins, from a file
# Under /admin/import/ rather than beside the thing they create: /admin/properties/{code} would
# read "template.csv" as the code of a hotel, and answer 404 about a route that exists.
@app.get("/admin/import/hotels.csv")
def properties_template(user: User = Depends(require_admin)):
    return Response(properties_import.template(), media_type="text/csv",
                    headers={"Content-Disposition": 'attachment; filename="hotels.csv"'})


@app.post("/admin/import/hotels", response_class=HTMLResponse)
async def properties_import_csv(request: Request, user: User = Depends(require_admin),
                                file: UploadFile = File(...)):
    """Add the hotels in one go.  Nothing is written unless every row is good."""
    store = require_db_store()
    raw = await file.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")            # a spreadsheet saved on a Windows machine
    parsed = properties_import.parse(text, existing_codes=store.properties(include_disabled=True))
    if not parsed.ok:
        return render(request, "admin_properties_import.html", errors=parsed.errors,
                      warnings=parsed.warnings, rows=[])
    for row in parsed.rows:
        try:
            store.save_property(**row.fields)
        except ValueError as e:                 # a mapping that no longer parses, say
            return render(request, "admin_properties_import.html", rows=[], warnings=[],
                          errors=[f"Line {row.line}: {e}. Nothing was imported."])
    state.reload_config()
    print(f"[properties] {user.username} imported {len(parsed.rows)} hotel(s) from {file.filename}")
    return render(request, "admin_properties_import.html", errors=[], warnings=parsed.warnings,
                  rows=parsed.rows)


# ------------------------------------------------------------------ logins, from a file
@app.get("/admin/import/logins.csv")
def users_template(user: User = Depends(require_admin)):
    """The file to fill in, carrying this deployment's own property codes in the examples."""
    body = users_import.template(sorted(state.props))
    return Response(body, media_type="text/csv",
                    headers={"Content-Disposition": 'attachment; filename="logins.csv"'})


@app.post("/admin/import/logins", response_class=HTMLResponse)
async def users_import_csv(request: Request, user: User = Depends(require_admin),
                           file: UploadFile = File(...), send: str = Form("yes")):
    """Create the logins in one go, and give each person a link to choose their own password.

    Nothing is written unless every row is good: half an import leaves somebody wondering
    which four of seven managers exist.
    """
    store = require_db_store()
    raw = await file.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")            # a spreadsheet saved on a Windows machine
    parsed = users_import.parse(text, known_properties=state.props,
                                existing_usernames=[u["username"] for u in store.all_users()],
                                protect=user.username)
    if not parsed.ok:
        return render(request, "admin_users_import.html", errors=parsed.errors, results=[],
                      can_email=mail.configured(), minutes=SET_PASSWORD_MINUTES)

    results = []
    for row in parsed.rows:
        store.save_user(username=row.username, role=row.role, properties=row.properties,
                        password_hash=(None if row.exists else unusable_password_hash()),
                        enabled=row.enabled, email=row.email)
        token = secrets.token_urlsafe(32)
        state.db.create_reset(row.username, token, SET_PASSWORD_MINUTES)
        link = f"{mail.base_url() or str(request.base_url).rstrip('/')}/reset?token={token}"
        sent = ""
        if send == "yes" and mail.configured():
            subject, body = mail.welcome_email(row.username, link, SET_PASSWORD_MINUTES,
                                               invited_by=user.username)
            ok, why = mail.send_reporting(row.email, subject, body)
            sent = "sent" if ok else f"not sent: {why.splitlines()[0]}"
        results.append({"row": row, "link": link, "sent": sent})
    state.reload_config()
    print(f"[users] {user.username} imported {len(results)} login(s) from {file.filename}")
    return render(request, "admin_users_import.html", errors=[], results=results,
                  can_email=mail.configured(), minutes=SET_PASSWORD_MINUTES)


# ------------------------------------------------------------------ worksheet import
def _pms_from_sheet(name: str) -> Optional[str]:
    words = [w for w in re.split(r"[^A-Za-z0-9]+", name) if w]
    for w in reversed(words):
        try:
            return get_parser(w).pms
        except KeyError:
            continue
    return None


@app.post("/admin/mapping/import/worksheet", response_class=HTMLResponse)
async def import_worksheet(request: Request, user: User = Depends(require_admin), file: UploadFile = File(...)):
    import io as _io
    import openpyxl
    try:
        wb = openpyxl.load_workbook(_io.BytesIO(await file.read()), data_only=True)
    except Exception as e:  # noqa: BLE001
        return RedirectResponse(f"/admin?err=Could+not+read+that+workbook:+{e}", status_code=303)
    applied, skipped = [], []
    for sheet in wb.worksheets:
        pms = _pms_from_sheet(sheet.title)
        if pms is None:
            skipped.append(f"{sheet.title}: not a PMS tab")
            continue
        head = {str(c.value).strip().lower(): c.column for c in sheet[4] if c.value}
        need = ("section", "pms code", "line as printed on the report", "your account code")
        if not all(k in head for k in need):
            skipped.append(f"{sheet.title}: unexpected columns")
            continue
        rules = []
        for row in range(5, sheet.max_row + 1):
            acct = sheet.cell(row, head["your account code"]).value
            label = sheet.cell(row, head["line as printed on the report"]).value
            if not acct or not label:
                continue
            rules.append(rule_for(str(label), str(sheet.cell(row, head["pms code"]).value or "").strip(),
                                  str(sheet.cell(row, head["section"]).value or "").strip(), str(acct).strip()))
        if not rules:
            skipped.append(f"{sheet.title}: no account codes filled in")
            continue
        targets = [c for c, p in state.props.items() if str(p.get("pms", "")).upper() == pms]
        if not targets:
            skipped.append(f"{sheet.title}: no property uses {pms}")
            continue
        for code in targets:
            prop = state.props[code]
            text = insert_rules(read_mapping_text(prop, state.store), rules,
                                f"# imported from the mapping worksheet {clock.today().isoformat()} by {user.username}")
            try:
                write_mapping_text(prop, state.store, text)
                applied.append(f"{code}: {len(rules)} rules")
            except ValueError as e:
                skipped.append(f"{code}: {e}")
    state.reload_config()
    return render(request, "admin_import_result.html", applied=applied, skipped=skipped)


@app.post("/runs/{run_id}/to-invoice")
def run_to_invoice(run_id: int, user: User = Depends(require_user)):
    _require_invoices()
    """The night-audit page decided this file is an invoice; move it across."""
    r = state.db.get_run(run_id)
    if not r:
        raise HTTPException(404)
    if not user.is_admin and r["uploaded_by"] != user.username:
        raise HTTPException(403)
    if r["status"] == "moved_to_invoice":
        raise HTTPException(400, "That file has already been moved to Invoices")
    res = RunResult.from_json(r["result_json"])
    if res.status != "looks_like_invoice":
        raise HTTPException(400, "That upload is not waiting to be moved")
    prop = res.property_code or (user.properties[0] if len(user.properties) == 1 else "")
    inv_id = _create_invoice(state.storage.local_path(r["stored_path"]), prop, user.username,
                             r["file_name"], r["stored_path"])
    state.db.mark(run_id, superseded=1, status="moved_to_invoice",
                  message=f"moved to invoice #{inv_id} by {user.username}")
    return RedirectResponse(f"/invoices/{inv_id}", status_code=303)


# ------------------------------------------------------------------ who has not reported
def _coverage_grid(days: int, end: Optional[date] = None) -> dict:
    """One row per property, one column per business date, for the last `days` nights.

    A night-audit pack covers the night before, so "last night" means yesterday's business
    date, uploaded this morning.  Three states per cell: reported and balanced, reported
    with a problem, or nothing at all.  Dates before a property's first-ever upload are
    left blank rather than flagged, so onboarding a hotel does not paint the page red.
    """
    end = end or clock.today() - timedelta(days=1)
    dates = [(end - timedelta(days=n)).isoformat() for n in range(days - 1, -1, -1)]
    runs = state.db.runs_between(dates[0], dates[-1])
    by_prop: dict[str, dict[str, dict]] = {}
    for r in runs:
        by_prop.setdefault(r["property_code"], {})[r["business_date"]] = r
    first_seen = state.db.first_business_date()
    rows, missing_last_night, problems_last_night = [], [], []
    for code, prop in state.props.items():
        cells, live_from = [], first_seen.get(code)
        for d in dates:
            run = by_prop.get(code, {}).get(d)
            if run is None:
                state_ = "before" if (live_from and d < live_from) or not live_from else "missing"
            else:
                state_ = "ok" if run["status"] == "ok" else "problem"
            cells.append({"date": d, "state": state_, "run": run})
        rows.append({"code": code, "name": prop.get("name", code), "pms": prop.get("pms", ""), "cells": cells})
        if cells[-1]["state"] == "missing":
            missing_last_night.append(code)
        elif cells[-1]["state"] == "problem":
            problems_last_night.append(code)
    return {"dates": dates, "rows": rows, "end": end,
            "missing_last_night": missing_last_night, "problems_last_night": problems_last_night}


@app.get("/missing", response_class=HTMLResponse)
def missing_report(request: Request, days: int = 14, day: Optional[str] = None,
                   user: User = Depends(require_admin)):
    days = max(1, min(days, 60))
    grid = _coverage_grid(days, date.fromisoformat(day) if day else None)
    return render(request, "missing.html", days=days, **grid)


@app.get("/missing.csv")
def missing_csv(days: int = 14, day: Optional[str] = None, user: User = Depends(require_admin)):
    import csv as _csv
    import io as _io
    days = max(1, min(days, 60))
    grid = _coverage_grid(days, date.fromisoformat(day) if day else None)
    buf = _io.StringIO(); w = _csv.writer(buf)
    w.writerow(["Property", "Name", "PMS", "Business date", "State", "Status", "Uploaded at", "Uploaded by"])
    for row in grid["rows"]:
        for cell in row["cells"]:
            r = cell["run"]
            w.writerow([row["code"], row["name"], row["pms"], cell["date"], cell["state"],
                        r["status"] if r else "", clock.show(r["created_at"]) if r else "",
                        r["uploaded_by"] if r else ""])
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="missing-uploads-{grid["end"]}.csv"'})


# ------------------------------------------------------------------ e-mail setup
@app.get("/admin/email", response_class=HTMLResponse)
def email_page(request: Request, user: User = Depends(require_admin), msg: str = "", err: str = "",
               note: str = ""):
    rows = []
    for key in mail.FORM_KEYS:
        secret = key in mail.SECRET_KEYS
        rows.append({"key": key, "source": mail.source_of(key), "secret": secret,
                     "help": mail.HELP.get(key, ""),
                     "value": "" if secret else mail.setting(key),
                     "is_set": bool(mail.setting(key)),
                     "locked": bool(os.environ.get(key))})
    return render(request, "admin_email.html", rows=rows, configured=mail.configured(),
                  msg=msg, err=err, note=note, checks=mail.checks(),
                  presets=mail.PRESETS, test_to=user.email or "")


@app.post("/admin/odoo-autopost")
def set_autopost(user: User = Depends(require_admin), autopost: str = Form("no")):
    if state.autopost_locked:
        return RedirectResponse("/admin?err=Pinned+on+the+host", status_code=303)
    state.db.save_settings({AUTOPOST_KEY: "yes" if autopost == "yes" else ""}, user.username)
    state.reload_config()
    return RedirectResponse("/admin?msg=" + ("Entries+will+be+posted+on+arrival" if state.autopost
                                             else "Entries+will+be+left+as+drafts"), status_code=303)


@app.post("/admin/email")
async def email_save(request: Request, user: User = Depends(require_admin)):
    form = await request.form()
    values = {}
    for key in mail.KEYS:
        if os.environ.get(key):
            continue                                   # pinned on the host, not ours to change
        v = (form.get(key) or "").strip()
        if key in mail.SECRET_KEYS and not v:
            continue                                   # blank means "leave the saved one alone"
        values[key] = v
    if form.get("clear_password"):
        values["SMTP_PASSWORD"] = ""
    state.db.save_settings(values, user.username)
    state.reload_config()
    return RedirectResponse("/admin/email?msg=Saved", status_code=303)


@app.post("/admin/email/preset")
def email_preset(user: User = Depends(require_admin), provider: str = Form("")):
    """Fill in the settings that are the same for everybody on a given service.

    Only the ones nobody chooses -- host, port, security, and the fixed login some services
    use.  The API key and the From address are the person's own and are left alone.
    """
    preset = mail.PRESETS.get(provider)
    if not preset:
        return RedirectResponse("/admin/email?err=Unknown+service", status_code=303)
    values = {k: v for k, v in preset.items()
              if k in mail.FORM_KEYS and v and not os.environ.get(k)}
    if values:
        state.db.save_settings(values, user.username)
        state.reload_config()
    return RedirectResponse(f"/admin/email?msg={quote_plus(preset['label'] + ' settings filled in')}"
                            f"&note={quote_plus(preset['note'])}", status_code=303)


@app.post("/admin/email/test")
def email_test(request: Request, to: str = Form(...), user: User = Depends(require_admin)):
    body = ("This is a test from the Night Audit portal.\n\n"
            "If you can read this, password resets and the daily missing-uploads notice will reach people.\n"
            f"Sent by {user.username}.\n")
    ok, why = mail.send_reporting(to.strip(), "Night Audit portal: test message", body)
    if ok:
        return RedirectResponse(
            f"/admin/email?msg={quote_plus('Test message sent to ' + to.strip())}", status_code=303)
    return RedirectResponse(f"/admin/email?err={quote_plus(why)}", status_code=303)
