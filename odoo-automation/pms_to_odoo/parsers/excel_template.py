"""Read the spreadsheet the GMs currently fill in by hand.

Harshil's email mentions an Excel template per property (example: a Hilton property).
Until the PMS parsers are finalised, this parser lets us feed Odoo from that sheet
directly, so the workflow can go live in two steps:

  step 1: GM keeps filling the spreadsheet -> we push it to Odoo automatically
  step 2: PMS report parsers replace the manual spreadsheet entirely

A cell map YAML describes where things live::

    sheet: Daily
    date_cell: B2
    property_cell: B1
    lines:                   # cell -> label, optional section
      B5:  {label: Room Revenue, section: revenue}
      B6:  {label: Occupancy Tax, section: tax}
      B12: {label: Visa/MC, section: settlement}
    ranges:                  # or a label column + amount column over a row range
      - {label_col: A, amount_col: B, first_row: 5, last_row: 40, section: ""}
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional

import openpyxl
import yaml

from ..models import DailyReport, ReportLine
from .base import BaseParser, parse_amount


class ExcelTemplateParser(BaseParser):
    pms = "XLSX"
    extensions = (".xlsx", ".xlsm")

    def __init__(self, cell_map_path: str | Path):
        self.cell_map = yaml.safe_load(Path(cell_map_path).read_text()) or {}

    def parse(self, path: str | Path, property_code: str,
              business_date: Optional[date] = None) -> DailyReport:
        wb = openpyxl.load_workbook(str(path), data_only=True)
        ws = wb[self.cell_map["sheet"]] if self.cell_map.get("sheet") else wb.active
        if business_date is None and self.cell_map.get("date_cell"):
            v = ws[self.cell_map["date_cell"]].value
            if isinstance(v, datetime):
                business_date = v.date()
            elif isinstance(v, date):
                business_date = v
            elif v:
                business_date = datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
        report = DailyReport(property_code=property_code, pms=self.pms,
                             business_date=business_date or date.today(), source_file=str(path))
        for cell, spec in (self.cell_map.get("lines") or {}).items():
            amt = _to_decimal(ws[cell].value)
            if amt is None:
                continue
            report.lines.append(ReportLine(spec["label"], amt, spec.get("section", ""),
                                           f"XLSX:{ws.title}!{cell}"))
        for rng in self.cell_map.get("ranges") or []:
            for row in range(int(rng["first_row"]), int(rng["last_row"]) + 1):
                label = ws[f"{rng['label_col']}{row}"].value
                amt = _to_decimal(ws[f"{rng['amount_col']}{row}"].value)
                if label and amt is not None:
                    report.lines.append(ReportLine(str(label).strip(), amt, rng.get("section", ""),
                                                   f"XLSX:{ws.title}!{rng['amount_col']}{row}"))
        return report


def _to_decimal(v) -> Optional[Decimal]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    return parse_amount(str(v))
